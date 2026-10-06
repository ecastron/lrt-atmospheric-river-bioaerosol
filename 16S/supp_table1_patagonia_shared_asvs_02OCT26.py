# supp_table1_patagonia_shared_asvs_02OCT26.py
# Supplementary Table 1. Writes the Patagonia-shared ASV table (summary sheet with criteria and groups, and one row
# per ASV with taxonomy, prevalence per pool, mean relative abundances and sequence) to Excel.
# Input  : 16S/figures/patagonia_transfer_sweep_02OCT26_supptable.tsv (patagonia_transfer_sweep_02OCT26.R)
# Output : supplementary_material/Supplementary_Table_1_patagonia_shared_ASVs.xlsx
# Run    : python3 16S/supp_table1_patagonia_shared_asvs_02OCT26.py

import csv, collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

rows = list(csv.DictReader(open('16S/figures/patagonia_transfer_sweep_02OCT26_supptable.tsv'), delimiter='\t'))
na = lambda v: '' if v in ('NA', None) else v
gen = lambda r: na(r['Genus']) or f"({na(r['Family']) or na(r['Order']) or na(r['Class']) or na(r['Phylum'])})"
n_yo = sum(r['group'] == 'Yelcho only' for r in rows)
n_ys = sum(r['group'] == 'Yelcho only' and int(r['n_yelcho_soil']) > 0 for r in rows)
order = ['New during/after AR', 'Baseline, increased during/after AR', 'Baseline, not increased', 'Yelcho only']

wb = Workbook()
ws = wb.active; ws.title = 'Summary'
ws['A1'] = ('Supplementary Table 1. Bacterial ASVs shared between Patagonian sources and Antarctic Peninsula air '
            'but absent or rare in local Antarctic soil, Southern Ocean air and negative controls.')
ws['A1'].font = Font(bold=True)
notes = [
    'Criteria: detected in at least 2 Patagonian air or 2 Patagonian bulk soil samples; detected in at least 2 '
    'Antarctic air samples (Risopatrón and Yelcho); present in no more than 5% of Risopatrón soil and rhizosphere samples '
    '(n = 63) and no more than 2% of Southern Ocean air samples over sea and over land (Malard et al. 2022, n = 153; the FEAST analysis in Figure 6 used the 135 over-sea samples with at least 1,000 reads); absent from all buffer '
    'controls (n = 5); not flagged by decontam (prevalence method, threshold 0.5). Chloroplast and mitochondrial ASVs '
    'were removed beforehand.',
    'Groups at Risopatrón: Baseline, detected in at least one of the 15 samples before the AR core (22 January to '
    '6 February 2022), split by whether mean relative abundance was higher During+After (7 to 14 February, n = 9); New, '
    'not detected before the AR and detected During+After; Yelcho only, not detected at Risopatrón (Yelcho sampling '
    f'began on 8 February). The local-soil filter used Risopatrón soils; {n_ys} of the {n_yo} Yelcho-only ASVs also occur in '
    'at least one of the 6 Yelcho soil samples (column "Yelcho soil") and may be of local origin.',
    'Relative abundances are means across samples, in percent of reads.']
for i, t in enumerate(notes, start=2):
    ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6); ws.row_dimensions[i].height = 75
hdr = ['Group', 'ASVs', 'Genera', 'Risopatrón air, Before (% of reads)', 'Risopatrón air, During+After (% of reads)', 'Most frequent genera (number of ASVs)']
r0 = len(notes) + 3
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=r0, column=j, value=h); c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True)
byg = collections.defaultdict(list)
for r in rows: byg[r['group']].append(r)
for k, g in enumerate(order, 1):
    L = byg[g]; cnt = collections.Counter(gen(r) for r in L)
    vals = [g, len(L), len({gen(r) for r in L if not gen(r).startswith('(')}),
            round(sum(float(r['ris_before_pct']) for r in L), 2), round(sum(float(r['ris_da_pct']) for r in L), 2),
            ', '.join(f'{n} ({v})' for n, v in cnt.most_common(8))]
    for j, v in enumerate(vals, 1):
        ws.cell(row=r0 + k, column=j, value=v).alignment = Alignment(wrap_text=True, vertical='top')
tot = r0 + len(order) + 1
ws.cell(row=tot, column=1, value='Total').font = Font(bold=True)
ws.cell(row=tot, column=2, value=len(rows))
ws.cell(row=tot, column=3, value=len({gen(r) for r in rows if not gen(r).startswith('(')}))
for col, w in zip('ABCDEF', [36, 8, 8, 20, 22, 70]): ws.column_dimensions[col].width = w

wa = wb.create_sheet('ASVs')
cols = [('asv', 'ASV'), ('group', 'Group'), ('patagonian_source', 'Patagonian source'), ('Phylum', 'Phylum'), ('Class', 'Class'),
        ('Order', 'Order'), ('Family', 'Family'), ('Genus', 'Genus'), ('Species', 'Species'),
        ('n_pat_air', 'Patagonian air (of 20)'), ('n_pat_soil', 'Patagonian bulk soil (of 19)'),
        ('n_ris_before', 'Risopatrón air Before (of 15)'), ('n_ris_da', 'Risopatrón air During+After (of 9)'),
        ('n_yelcho', 'Yelcho air (of 30)'), ('n_local_soil', 'Risopatrón soil and rhizosphere (of 63)'), ('n_yelcho_soil', 'Yelcho soil (of 6)'),
        ('n_so_air', 'Southern Ocean air (of 153)'), ('ris_before_pct', 'Risopatrón air Before, mean %'),
        ('ris_da_pct', 'Risopatrón air During+After, mean %'), ('ant_air_mean_pct', 'Antarctic air, mean %'), ('sequence', 'Sequence')]
for j, (_, h) in enumerate(cols, 1):
    c = wa.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True)
for i, r in enumerate(rows, 2):
    for j, (k, _) in enumerate(cols, 1):
        v = na(r.get(k, ''))
        if k.startswith('n_'): v = int(v)
        elif k.endswith('_pct'): v = round(float(v), 4)
        wa.cell(row=i, column=j, value=v)
wa.freeze_panes = 'B2'
for j in range(1, len(cols) + 1): wa.column_dimensions[get_column_letter(j)].width = 14
wa.column_dimensions['B'].width = 32; wa.column_dimensions[get_column_letter(len(cols))].width = 60
wb.save('supplementary_material/Supplementary_Table_1_patagonia_shared_ASVs.xlsx')
print('saved', len(rows), 'ASVs')
